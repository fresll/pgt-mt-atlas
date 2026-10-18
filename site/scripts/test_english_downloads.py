import csv
import hashlib
from pathlib import Path
import tempfile
import unittest
import zipfile

import openpyxl

from build_english_downloads import ROOT, CYRILLIC, build_csv, build_template


class EnglishDownloads(unittest.TestCase):
    def test_csv_preserves_schema_ids_numeric_values_and_source_files(self):
        source = ROOT / 'public/data/csv'
        hashes = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in source.glob('*.csv')}
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            self.assertEqual(build_csv(source, output), len(hashes))
            for name in hashes:
                with (source / name).open(encoding='utf-8-sig', newline='') as stream:
                    before = list(csv.reader(stream))
                with (output / name).open(encoding='utf-8-sig', newline='') as stream:
                    after = list(csv.reader(stream))
                self.assertEqual(before[0], after[0])
                self.assertEqual(len(before), len(after))
                for original, translated in zip(before, after, strict=True):
                    self.assertEqual(len(original), len(translated))
                    for old, new in zip(original, translated, strict=True):
                        self.assertFalse(CYRILLIC.search(new), new)
                        if not CYRILLIC.search(old):
                            self.assertEqual(old, new)
                self.assertEqual(hashes[name], hashlib.sha256((source / name).read_bytes()).hexdigest())

    def test_workbook_keeps_fields_constraints_and_english_definitions(self):
        source = ROOT / 'public/templates/PGT-mt_Research_Data_Submission_v1.0.xlsx'
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        original = openpyxl.load_workbook(source)
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / source.name
            build_template(source, target)
            translated = openpyxl.load_workbook(target)
            self.assertEqual(original.sheetnames, translated.sheetnames)
            for name in ['Study', 'Cycle', 'Embryos', 'Samples_Measurements', 'Methods', 'Clinical_Outcomes', 'Lists']:
                before, after = original[name], translated[name]
                self.assertEqual([c.value for c in before[1]], [c.value for c in after[1]])
                for a, b in zip(before.data_validations.dataValidation, after.data_validations.dataValidation, strict=True):
                    for attr in ['sqref', 'type', 'operator', 'formula1', 'formula2', 'allowBlank']:
                        self.assertEqual(getattr(a, attr), getattr(b, attr), (name, attr))
                self.assertEqual(list(before.values)[2:], list(after.values)[2:])
            source_definitions = [(r[0], r[1], r[2], r[8], r[10]) for r in list(original['Data_Dictionary'].values)[1:]]
            target_definitions = [(r[0], r[1], r[2], r[7], r[8]) for r in list(translated['Data_Dictionary'].values)[1:]]
            self.assertEqual(source_definitions, target_definitions)
            with zipfile.ZipFile(target) as archive:
                for name in archive.namelist():
                    if name.endswith('.xml'):
                        self.assertFalse(CYRILLIC.search(archive.read(name).decode()), name)
        self.assertEqual(digest, hashlib.sha256(source.read_bytes()).hexdigest())


if __name__ == '__main__':
    unittest.main()
