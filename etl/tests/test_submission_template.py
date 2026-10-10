"""Шаблон PGT-mt Research Data Submission: структура, связи, различие «0 / пусто / ниже предела», отсутствие идентификаторов."""
import io
import json
import re

import openpyxl
import pytest

from atlas_etl import submission_template as st

DATA_SHEETS = ['Study', 'Cycle', 'Embryos', 'Samples_Measurements', 'Methods', 'Clinical_Outcomes']


@pytest.fixture(scope='module')
def wb():
    return openpyxl.load_workbook(io.BytesIO(st.workbook_bytes()))


def test_committed_files_are_up_to_date():
    assert st.main(['--check']) == 0


def test_build_is_deterministic():
    assert st.workbook_bytes() == st.workbook_bytes()


def test_sheets(wb):
    assert wb.sheetnames == ['Instructions', *DATA_SHEETS, 'Data_Dictionary', 'Lists']


def test_headers_match_dictionary(wb):
    dd = list(wb['Data_Dictionary'].iter_rows(min_row=2, values_only=True))
    for sh in st.SHEETS:
        header = [c.value for c in wb[sh.name][1] if c.value]
        assert header == [f.code for f in sh.fields]
        assert len(header) == len(set(header)), sh.name
        assert [r[1] for r in dd if r[0] == sh.name] == header
    # у каждого поля есть определения на обоих языках
    assert all(r[8] and r[9] for r in dd)


def test_keys_link_sheets():
    codes = {sh.name: {f.code: f for f in sh.fields} for sh in st.SHEETS}
    assert codes['Study']['study_id'].type == 'id'
    for sheet, fields in codes.items():
        for f in fields.values():
            if f.type != 'ref':
                continue
            for target in f.allowed.split(' | '):
                tsheet, tfield = target.split('.')
                assert codes[tsheet][tfield].type == 'id', (sheet, f.code, target)


def test_enums_have_lists_and_validation(wb):
    for sh in st.SHEETS:
        ws = wb[sh.name]
        validated = {str(r).split(':')[0].rstrip('0123456789') for dv in ws.data_validations.dataValidation for r in dv.sqref.ranges}
        for j, f in enumerate(sh.fields, start=1):
            if f.type == 'enum':
                assert f.allowed in st.LISTS, (sh.name, f.code)
            if f.type in ('enum', 'percent', 'integer', 'decimal'):
                assert openpyxl.utils.get_column_letter(j) in validated, (sh.name, f.code)


def test_zero_missing_and_lod_are_distinguished():
    q = st.LISTS['value_qualifier']
    for code in ('exact', 'below_lod', 'interval', 'less_than', 'greater_than', 'failed_qc', 'not_measured', 'not_reported'):
        assert code in q
    sm = {f.code: f for f in st.SHEETS[3].fields}
    assert sm['value_qualifier'].req == 'R'
    assert 'lod_pct' in sm and 'value_low_pct' in sm and 'value_high_pct' in sm
    text = ' '.join(t for _, t in st.INSTRUCTIONS_EN)
    assert 'Never type 0 for unknown' in text and 'below_lod' in text


def test_complete_cycles_and_negative_outcomes_are_captured():
    cyc = {f.code for f in st.SHEETS[1].fields}
    assert {'oocytes_retrieved', 'oocytes_mii', 'zygotes_2pn', 'blastocysts_total', 'embryos_biopsied', 'embryos_suitable', 'cycle_endpoint'} <= cyc
    assert 'no_suitable_embryo' in st.LISTS['cycle_endpoint']
    emb = {f.code for f in st.SHEETS[2].fields}
    assert {'stage_at_biopsy', 'morphology_grade', 'pgta_result', 'transfer_status', 'reason_not_transferred'} <= emb
    for ev in ('no_suitable_embryo', 'transfer_no_implantation', 'miscarriage', 'live_birth', 'follow_up_visit'):
        assert ev in st.LISTS['outcome_event']


def test_evidence_tiers_and_permissions():
    assert st.LISTS['evidence_tier'] == ['peer_reviewed_published', 'published_not_peer_reviewed', 'unpublished_research_data']
    study = {f.code: f for f in st.SHEETS[0].fields}
    for k in ('ethics_approval', 'consent_for_research_use', 'deidentification_confirmed', 'authority_to_share', 'permission_public_display', 'overlap_with_published'):
        assert study[k].req == 'R', k


def test_no_direct_identifier_fields():
    banned = re.compile(r'(^|_)(name|surname|initials|dob|birth_date|date_of_birth|address|phone|mrn|record_number|passport|email)($|_)')
    for sh in st.SHEETS[1:]:  # Study содержит только профессиональные сведения об исследователях
        for f in sh.fields:
            assert not banned.search(f.code), (sh.name, f.code)
            assert f.type != 'date', (sh.name, f.code)
    # возраст — только группами
    assert {f.code: f for f in st.SHEETS[1].fields}['age_group'].type == 'enum'


def test_summary_json():
    s = json.loads(st.OUT_JSON.read_text(encoding='utf-8'))
    assert s['file'] == f'templates/{st.FILENAME}'
    assert [x['name'] for x in s['sheets']] == DATA_SHEETS
