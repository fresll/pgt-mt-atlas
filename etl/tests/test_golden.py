"""Эталонные тесты: ETL обязан воспроизводить числа рукописи обзора (docs/data_inventory.md, раздел 3)."""
import json
from fractions import Fraction

import pytest

from atlas_etl import build as b


@pytest.fixture(scope='module')
def atlas():
    b.check_manifest(b.sha256(b.SOURCE_XLSX))
    return b.build()


def test_fig3_counts(atlas):
    s = atlas['stats']['fig3']
    assert s['n'] == 187
    expected = {  # вариант: (n, ≤18%, ≤30%)
        'm.3243A>G': (87, 41, 52), 'm.14487T>C': (28, 6, 8), 'm.8993T>G': (26, 24, 24),
        'm.3697G>A': (21, 1, 3), 'm.10197G>A': (9, 5, 5), 'm.10191T>C': (8, 8, 8), 'm.3250T>C': (8, 8, 8),
    }
    got = {v: (x['n'], x['le18']['k'], x['le30']['k']) for v, x in s['by_variant'].items()}
    assert got == expected
    assert s['pooled_le18'] == {'k': 93, 'n': 187}


def test_fig3_intervals_not_imputed(atlas):
    iv = [m for m in atlas['fig3'] if m['q'] == 'interval']
    assert [(m['low'], m['high'], m['value']) for m in iv] == [(26, 45, None), (84, 88, None)]


def test_fig4(atlas):
    s = atlas['stats']['fig4']
    assert (s['n_embryos'], s['n_pathogenic'], s['n_hv2']) == (40, 21, 19)
    # исходная версия xlsx воспроизводит рукопись (2,1); после корректировок по ESM 2 (A3) медиана пересчитана: 2,2
    assert round(atlas['stats']['fig4_source_xlsx']['median_range_pathogenic'], 1) == 2.1
    assert round(s['median_range_pathogenic'], 1) == 2.2
    assert s['median_range_hv2'] == 1.0
    assert round(s['median_range_all'], 1) == 1.2
    assert round(s['share_le5_pathogenic'], 2) == 0.76 and s['share_le5_all'] == 0.85
    assert s['max_range_pathogenic'] == 8
    assert (s['n_cleavage'], s['median_range_cleavage'], s['max_range_cleavage']) == (22, 1, 6)


def test_fig5(atlas):
    s = atlas['stats']['fig5']
    assert (s['n_pairs'], s['n_children'], s['n_low_embryo'], s['n_concordant_low']) == (22, 8, 20, 16)
    assert s['discordant_ids'] == ['C5e', 'C5f', 'C5g', 'C5h']
    excluded = {p['id']: p['q'] for p in atlas['fig5'] if not p['in_figure']}
    assert excluded == {'C5d': 'less_than', 'C13': 'qualitative'}


def test_fig6_example():
    """Пример из рукописи: m.3243A>G, 35 лет → ≈24 / 55 / 80% при N = 1 / 3 / 6."""
    p = Fraction(41, 87) * Fraction(1, 2)
    got = [round(float(1 - (1 - p) ** n), 2) for n in (1, 3, 6)]
    assert got == [0.24, 0.55, 0.8]


def test_no_invented_doi(atlas):
    pubs = {p['id']: p for p in atlas['publications']}
    assert pubs['thorburn2009']['doi'] is None
    for m in atlas['fig3']:
        assert m['doi_in_sheet'] is None or m['doi_in_sheet'] == pubs[m['pub']]['doi']


def test_tofilo_published(atlas):
    t = {p['id']: p for p in atlas['publications']}['tofilo2026']
    assert t['pub_status'] == 'published'
    assert t['doi'] == '10.1007/s10815-026-04012-y'
    assert t['published_date'] == '2026-08-27' and t['pmid'] == '42649365'


def test_tofilo_esm2_verification(atlas):
    v = atlas['verification']
    status = {(r['embryo'], r['esm_fraction']): r['status'] for r in v if r['embryo']}
    assert status[('Tofilo_6', 'icm')] == 'value_mismatch'
    assert status[('Tofilo_6', 'remaining_te')] == 'value_mismatch'
    assert status[('Tofilo_3', 'combined_te_icm')] == 'fraction_label_mismatch'
    assert status[('Tofilo_4', 'combined_te_icm')] == 'fraction_label_mismatch'
    assert sum(s == 'match' for s in status.values()) == 12


def test_values_in_range(atlas):
    for m in atlas['fig3'] + atlas['fig4']:
        if m.get('value') is not None:
            assert 0 <= m['value'] <= 100


def test_json_serializable(atlas):
    json.dumps(atlas, ensure_ascii=False)


# --- ESM 2 Tofilo 2026 обязателен: без него сборка не должна молча снимать пометки расхождений

def test_missing_esm2_fails_build(tmp_path, monkeypatch):
    monkeypatch.setattr(b, 'TOFILO_ESM2', tmp_path / 'absent.xlsx')
    with pytest.raises(b.EtlError, match='10815_2026_4012_MOESM2_ESM.xlsx'):
        b.build()
    assert b.main([]) == 1  # обычный запуск ETL завершается ошибкой, файлы не пишутся
    assert b.main(['--check']) == 1


def test_unreadable_esm2_fails(atlas, tmp_path):
    bad = tmp_path / 'esm.xlsx'
    bad.write_bytes(b'not an xlsx')
    with pytest.raises(b.EtlError, match='не удалось прочитать'):
        b.verify_tofilo(atlas['fig3'], atlas['fig4'], bad)


def test_esm2_wrong_structure_fails(atlas, tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(['Embryo', 'Fraction', 'Value'])
    wb.active.append(['Embryo 1', 'TE biopsy', 11.4])
    path = tmp_path / 'esm.xlsx'
    wb.save(path)
    with pytest.raises(b.EtlError, match='неожиданная структура'):
        b.verify_tofilo(atlas['fig3'], atlas['fig4'], path)


def test_esm2_discrepancy_detected_not_rewritten(atlas, tmp_path):
    """Подмена значения в копии ESM выявляется как расхождение; значения набора остаются прежними."""
    import openpyxl
    src = openpyxl.load_workbook(b.TOFILO_ESM2)
    ws = src.active
    row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == 'Embryo 1' and r[1].value == 'ICM')
    original = row[8].value
    row[8].value = original + 5
    path = tmp_path / 'esm.xlsx'
    src.save(path)
    before = [(f['embryo'], f['fraction'], f['value']) for f in atlas['fig4']]
    v = {(r['embryo'], r['esm_fraction']): r for r in b.verify_tofilo(atlas['fig3'], atlas['fig4'], path) if r['embryo']}
    assert v[('Tofilo_1', 'icm')]['status'] == 'value_mismatch'
    assert v[('Tofilo_1', 'icm')]['dataset_value'] == original
    assert [(f['embryo'], f['fraction'], f['value']) for f in atlas['fig4']] == before


# --- A3: прозрачный слой корректировок Tofilo 2026 по опубликованному ESM 2

def _fig4(atlas, embryo):
    return {f['fraction']: f for f in atlas['fig4'] if f['embryo'] == embryo}


def test_tofilo_corrections_follow_esm2_with_provenance(atlas):
    e6 = _fig4(atlas, 'Tofilo_6')
    assert (e6['icm']['value'], e6['icm']['value_source']) == (11.3, 11.4)
    assert (e6['remaining_te']['value'], e6['remaining_te']['value_source']) == (11.7, 11.9)
    assert e6['icm']['corrections'] == ['corr-tofilo2026-e6-icm']
    for emb, value in (('Tofilo_3', 13.5), ('Tofilo_4', 90.0)):
        f = _fig4(atlas, emb)
        assert 'remaining_te' not in f
        assert f['combined_te_icm']['value'] == value
        assert f['combined_te_icm']['fraction_source'] == 'remaining_te'
    assert {c['id'] for c in atlas['corrections']} == {
        'corr-tofilo2026-e6-icm', 'corr-tofilo2026-e6-remaining-te', 'corr-tofilo2026-e3-fraction', 'corr-tofilo2026-e4-fraction'}
    for c in atlas['corrections']:
        assert c['source']['locator'] and c['source']['publication'] == 'tofilo2026'


def test_corrections_do_not_touch_source_workbook():
    b.check_manifest(b.sha256(b.SOURCE_XLSX))


def test_correction_stops_build_if_source_changed(atlas):
    rec = [dict(f, src=dict(f['src'])) for f in atlas['fig4']]
    for f in rec:  # вернуть исходные значения xlsx
        f['value'] = f.get('value_source', f['value'])
        f['fraction'] = f.get('fraction_source', f['fraction'])
        f.pop('value_source', None), f.pop('fraction_source', None), f.pop('corrections', None)
    corr = b.load_yaml('curation/corrections.yaml')['corrections']
    bad = [dict(corr[0], **{'from': 11.5})]
    with pytest.raises(b.EtlError, match='Исходник изменился'):
        b.apply_corrections(rec, bad)


def test_verification_reports_corrected_discrepancies(atlas):
    v = {(r['embryo'], r['esm_fraction']): r for r in atlas['verification'] if r['embryo']}
    for key in (('Tofilo_6', 'icm'), ('Tofilo_6', 'remaining_te'), ('Tofilo_3', 'combined_te_icm'), ('Tofilo_4', 'combined_te_icm')):
        assert v[key]['resolution'] == 'corrected', key
    assert v[('Tofilo_6', 'icm')]['atlas_value'] == 11.3
    assert v[('Tofilo_3', 'combined_te_icm')]['atlas_fraction'] == 'combined_te_icm'
    assert v[('Tofilo_1', 'icm')]['resolution'] is None


# --- B2: стабильные идентификаторы и версии набора

def test_ids_do_not_depend_on_row_position(atlas):
    reg = b.load_registry()
    shuffled = [dict(m, src=dict(m['src'], row=1000 - m['src']['row'])) for m in reversed(atlas['fig3'])]
    b.assign_ids({'fig3': shuffled}, {'records': [dict(e) for e in reg['records']]}, False, '2026-10-11')
    unique = {}
    for m in atlas['fig3']:
        unique.setdefault(json.dumps(b.fingerprint('fig3', m), sort_keys=True), []).append(m['id'])
    for m in shuffled:
        ids = unique[json.dumps(b.fingerprint('fig3', m), sort_keys=True)]
        assert m['id'] in ids
        if len(ids) == 1:
            assert m['id'] == ids[0]


def test_new_record_needs_curated_id(atlas):
    reg = b.load_registry()
    new = dict(atlas['fig3'][0], value=55.5, src=dict(atlas['fig3'][0]['src'], row=999))
    recs = [dict(m) for m in atlas['fig3']] + [new]
    with pytest.raises(b.EtlError, match='--assign-ids'):
        b.assign_ids({'fig3': recs}, {'records': [dict(e) for e in reg['records']]}, False, '2026-10-11')
    recs = [dict(m) for m in atlas['fig3']] + [dict(new)]
    _, new_ids = b.assign_ids({'fig3': recs}, {'records': [dict(e) for e in reg['records']]}, True, '2026-10-11')
    assert len(new_ids) == 1 and new_ids[0].startswith(f"HL-{new['pub']}-")
    assert [m['id'] for m in recs[:-1]] == [m['id'] for m in atlas['fig3']]


def test_legacy_ids_are_kept_as_aliases(atlas):
    hu = [m for m in atlas['fig3'] if m['pub'] == 'hu2026' and m['value'] == 17.8]
    assert len(hu) == 1 and hu[0]['legacy_id'] == 'f3-183' and hu[0]['id'].startswith('HL-hu2026-')
    ids = [m['id'] for m in atlas['fig3']] + [f['id'] for f in atlas['fig4']] + [c['id'] for c in atlas['cleavage']]
    assert len(ids) == len(set(ids)) == 300


def test_baseline_release_is_reproduced(atlas):
    meta = atlas['meta']
    assert meta['baseline']['version'] == '2026-10-10'
    assert meta['baseline']['counts']['landscape_values'] == 187
    assert meta['current_counts'] == meta['baseline']['counts']


# --- раздел 7: основания классификации, роли источников; A4–A5: модель и ориентиры

def test_variant_classification_has_basis(atlas):
    vc = atlas['variant_classes']
    assert vc['m.3250T>C']['class'] == 'uncertain_significance'
    assert vc['m.3250T>C']['status'] == 'stated_in_source' and vc['m.3250T>C']['publication'] == 'hu2026'
    assert 'uncertain significance' in vc['m.3250T>C']['locator']
    for v, e in vc.items():
        assert e['checked'] and (e['locator'] or e['status'] == 'not_verified'), v


def test_every_publication_has_a_role(atlas):
    for p in atlas['publications']:
        assert p['roles'], p['id']
    roles = {p['id']: {r['role'] for r in p['roles']} for p in atlas['publications']}
    assert roles['hellebrekers2012'] == {'threshold'} and roles['volodyaev2025'] == {'model_parameter'}


def test_thresholds_are_not_universal(atlas):
    th = {t['id']: t for t in atlas['thresholds']}
    assert th['t18']['role'] == 'study_estimate' and 'мышеч' in th['t18']['ru']
    assert th['t30']['role'] == 'programme_criterion' and 'Newcastle' in th['t30']['ru']
    for k in ('lod1', 'high70'):
        assert th[k]['role'] == 'chart_convention' and th[k]['publication'] is None
    for t in th.values():
        assert t['checked']


def test_model_p_eupl_is_labelled_scenario(atlas):
    m = atlas['model']
    assert [x['value'] for x in m['p_eupl']] == [0.7, 0.5, 0.3]  # параметры не менялись
    assert m['p_eupl_status'] == 'scenario'
    assert 'xfnr.2024.100082' in m['p_eupl_locator']
