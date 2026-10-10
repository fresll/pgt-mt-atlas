"""Шаблон PGT-mt Research Data Submission (xlsx) — единый источник описания полей.

Шаблон — самостоятельный научный стандарт Atlas для предложения данных ПГТ-мт; научной основой
служит минимальный набор данных (пациентка и цикл, эмбрион, метод, исход), дополненный полями,
которые нужны для учёта полных циклов и отрицательных результатов.

    python3 -m atlas_etl.submission_template            # из каталога etl/
    python3 -m atlas_etl.submission_template --check    # сверить с сохранёнными файлами

Пишет:
    site/public/templates/<FILENAME>          — скачиваемый шаблон
    site/src/data/submissionTemplate.json     — сводка для страницы Contribute Data

Файл детерминирован (фиксированные даты в свойствах и в zip), поэтому повторная сборка
не меняет его без изменения описания полей.
"""
from __future__ import annotations

import argparse
import datetime as dt
import io
import json
import re
import sys
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
VERSION = '1.0'
RELEASED = '2026-10-10'
FILENAME = f'PGT-mt_Research_Data_Submission_v{VERSION}.xlsx'
OUT_XLSX = ROOT / 'site' / 'public' / 'templates' / FILENAME
OUT_JSON = ROOT / 'site' / 'src' / 'data' / 'submissionTemplate.json'
DATA_ROWS = 500  # строки с проверкой значений в каждом листе данных
FIRST_DATA_ROW = 3

# Scientific Sketchbook: синие чернила, красный акцент, белая бумага
INK = '1F3FB4'
INK2 = '34406E'
MUTED = '5F6787'
RED = 'D42A2F'
RULE = 'CFD6EF'

# ---------------------------------------------------------------- словари допустимых значений

LISTS: dict[str, list[str]] = {
    'yes_no_unknown': ['yes', 'no', 'unknown'],
    'yes_no': ['yes', 'no'],
    'evidence_tier': [
        'peer_reviewed_published',      # рецензированная публикация
        'published_not_peer_reviewed',  # препринт, тезисы, диссертация, отчёт
        'unpublished_research_data',    # неопубликованные данные исследовательской группы
    ],
    'submission_scope': ['complete_cycles', 'partial_cycles', 'embryo_level_only', 'outcome_follow_up_only', 'methods_validation'],
    'overlap': ['no_overlap', 'partial_overlap', 'full_overlap_update', 'unknown'],
    'age_group': ['<30', '30-34', '35-37', '38-40', '41-42', '>42', 'unknown'],
    'family_history': ['affected_child', 'affected_relatives', 'affected_patient', 'carrier_identified_by_testing', 'apparently_de_novo', 'unknown'],
    'carrier_status': ['asymptomatic', 'mildly_symptomatic', 'symptomatic', 'unknown'],
    'variant_class': ['pathogenic', 'likely_pathogenic', 'uncertain_significance', 'neutral_polymorphism_control'],
    'fertilisation_method': ['ICSI', 'conventional_IVF', 'mixed', 'unknown'],
    'cycle_endpoint': [
        'transfer_performed',
        'no_suitable_embryo',
        'no_embryo_reached_biopsy',
        'no_oocytes_or_no_fertilisation',
        'all_embryos_cryopreserved_pending',
        'cycle_cancelled',
        'research_only_no_transfer',
        'unknown',
    ],
    'embryo_origin': ['clinical_pgt_cycle', 'donated_for_research', 'other'],
    'stage': ['zygote', 'cleavage_d2_d3', 'morula_d4', 'blastocyst_d5', 'blastocyst_d6', 'blastocyst_d7', 'arrested', 'degenerated', 'unknown'],
    'morphology_system': ['Gardner', 'Istanbul_consensus_2011', 'ESHRE_ALPHA_2025', 'time_lapse_score', 'clinic_specific', 'not_assessed'],
    'pgta_result': ['euploid', 'aneuploid', 'mosaic_low_level', 'mosaic_high_level', 'segmental_only', 'no_result', 'not_tested'],
    'suitability': ['suitable', 'not_suitable', 'undetermined'],
    'transfer_status': ['transferred_fresh', 'transferred_frozen', 'cryopreserved_not_transferred', 'discarded', 'donated_to_research', 'arrested_before_decision', 'unknown'],
    'not_transferred_reason': [
        'heteroplasmy_above_cutoff',
        'aneuploid_or_mosaic',
        'poor_morphology_or_arrest',
        'no_pgt_result',
        'patient_decision',
        'surplus_after_successful_transfer',
        'research_protocol',
        'other',
        'not_applicable',
    ],
    'subject_level': ['carrier_patient', 'oocyte', 'polar_body', 'embryo', 'fetus', 'child'],
    'sample_type': [
        'blood', 'urine_sediment', 'buccal_epithelium', 'skeletal_muscle', 'hair_follicles', 'skin_fibroblasts',
        'oocyte', 'polar_body_1', 'polar_body_2', 'single_blastomere', 'multiple_blastomeres',
        'te_biopsy', 'icm', 'remaining_te', 'whole_embryo', 'embryo_half', 'spent_culture_medium',
        'chorionic_villi', 'amniocytes', 'placenta', 'cord_blood', 'fetal_tissue', 'other',
    ],
    'timepoint_unit': ['embryo_day', 'gestational_week', 'child_age_months', 'adult_age_band'],
    'measurement_kind': ['single_replicate', 'mean_of_replicates', 'single_measurement'],
    'value_qualifier': [
        'exact',          # измеренное число (в т. ч. 0, если метод действительно выдал количественный ноль)
        'below_lod',      # ниже предела обнаружения: value пусто, lod_pct обязателен
        'less_than',      # опубликовано только «< X»: X в value_high_pct
        'greater_than',   # опубликовано только «> X»: X в value_low_pct
        'interval',       # опубликован только интервал: value_low_pct и value_high_pct
        'qualitative',    # только качественный результат (text_result)
        'failed_qc',      # анализ выполнен, но не прошёл контроль качества
        'not_measured',   # образец получен, анализ не проводился
        'not_reported',   # измерено, но значение недоступно
    ],
    'analytical_method': [
        'NGS_amplicon', 'NGS_whole_mtDNA', 'NGS_whole_genome_low_pass', 'ddPCR', 'qPCR_ARMS', 'PCR_RFLP',
        'pyrosequencing', 'fluorescent_PCR_fragment', 'Sanger_quantitative', 'SNaPshot', 'other',
    ],
    'wga': ['none', 'MDA', 'MALBAC', 'PicoPLEX_SurePlex', 'PEP_or_DOP', 'other', 'unknown'],
    'outcome_event': [
        'no_suitable_embryo',
        'no_transfer_other_reason',
        'transfer_no_implantation',
        'biochemical_pregnancy',
        'clinical_pregnancy_ongoing',
        'miscarriage',
        'ectopic_pregnancy',
        'termination_after_prenatal_diagnosis',
        'termination_other_reason',
        'stillbirth',
        'live_birth',
        'follow_up_visit',
        'unknown',
    ],
    'child_status': ['no_signs_of_mitochondrial_disease', 'signs_or_symptoms', 'diagnosed_mitochondrial_disease', 'not_assessed', 'lost_to_follow_up'],
    'prenatal_test': ['none', 'CVS', 'amniocentesis', 'both', 'declined', 'unknown'],
}

# что Atlas сохраняет в пустой ячейке; используется в инструкциях и словаре
BLANK_RULE_EN = 'Blank = not reported / unknown. 0 = counted and none. Never type 0 for “unknown”.'
BLANK_RULE_RU = 'Пусто = не сообщается / неизвестно. 0 = посчитано, и таких нет. Не ставьте 0 вместо «неизвестно».'


@dataclass
class F:
    code: str
    label_en: str
    label_ru: str
    desc_en: str
    desc_ru: str
    type: str = 'text'          # id, ref, text, integer, percent, decimal, enum, year
    req: str = 'O'              # R — обязательное, C — обязательно при условии, O — по возможности
    unit: str = ''
    allowed: str = ''           # имя словаря LISTS для enum или текстовое правило
    example: str = ''


@dataclass
class Sheet:
    name: str
    title_en: str
    title_ru: str
    purpose_en: str
    purpose_ru: str
    fields: list[F] = field(default_factory=list)


def _count(code, en, ru, den, dru, req='R', ex=''):
    return F(code, en, ru, den + ' ' + BLANK_RULE_EN, dru + ' ' + BLANK_RULE_RU, 'integer', req, 'count', '≥0', ex)


SHEETS: list[Sheet] = [
    Sheet(
        'Study', 'Study', 'Исследование',
        'One row per study or data set: who generated the data, where, and in what publication status.',
        'Одна строка на исследование или набор данных: кто, где и в каком статусе публикации получил данные.',
        [
            F('study_id', 'Study ID', 'ID исследования', 'Short identifier chosen by the submitter; used as the key in all other sheets.', 'Краткий идентификатор, выбранный автором; ключ для всех остальных листов.', 'id', 'R', example='LAB1-2026'),
            F('study_title', 'Title', 'Название', 'Title of the publication or a descriptive title of the unpublished data set.', 'Название публикации или описательное название неопубликованного набора.', 'text', 'R', example='PGT-mt cycles for m.3243A>G, 2019–2025'),
            F('institution', 'Institution', 'Организация', 'Clinic, laboratory or university that generated the data.', 'Клиника, лаборатория или университет, получившие данные.', 'text', 'R'),
            F('country', 'Country', 'Страна', 'Country where the cycles were performed.', 'Страна, где проводились циклы.', 'text', 'R'),
            F('corresponding_researcher', 'Corresponding researcher', 'Ответственный исследователь', 'Researcher responsible for the submission (not a patient). Professional contact only.', 'Исследователь, отвечающий за передачу данных (не пациент). Только профессиональный контакт.', 'text', 'R'),
            F('researcher_orcid', 'ORCID', 'ORCID', 'ORCID iD of the corresponding researcher, if available.', 'ORCID ответственного исследователя, если есть.', 'text', 'O', example='0000-0000-0000-0000'),
            F('study_authors', 'Authors / data generators', 'Авторы / создатели данных', 'All authors of the publication, or the people who generated unpublished data, in the order they should be credited.', 'Все авторы публикации или люди, получившие неопубликованные данные, в порядке, в котором их следует указывать.', 'text', 'R'),
            F('evidence_tier', 'Publication status', 'Статус публикации', 'Peer-reviewed publication; published without independent peer review (preprint, abstract, thesis, report); or unpublished research data.', 'Рецензированная публикация; опубликовано без независимого рецензирования (препринт, тезисы, диссертация, отчёт) или неопубликованные данные.', 'enum', 'R', allowed='evidence_tier'),
            F('doi', 'DOI', 'DOI', 'DOI of the publication; required unless the data are unpublished.', 'DOI публикации; обязателен, кроме неопубликованных данных.', 'text', 'C', example='10.xxxx/xxxxx'),
            F('pmid', 'PMID', 'PMID', 'PubMed ID, if indexed.', 'PubMed ID, если есть.', 'text', 'O'),
            F('citation', 'Full citation', 'Полная ссылка', 'Full reference of the publication, preprint or abstract.', 'Полная ссылка на публикацию, препринт или тезисы.', 'text', 'C'),
            F('source_location', 'Where in the source', 'Где в источнике', 'For published data: table, figure or supplementary file that holds the values.', 'Для опубликованных данных: таблица, рисунок или приложение, где приведены значения.', 'text', 'C', example='Supplementary Table S2'),
            F('unpublished_status', 'Status of unpublished data', 'Статус неопубликованных данных', 'For unpublished data: e.g. “manuscript in preparation”, “submitted to <journal>”, “not planned for publication”.', 'Для неопубликованных данных: например, «рукопись готовится», «направлена в <журнал>», «публикация не планируется».', 'text', 'C'),
            F('study_period', 'Period of cycles', 'Период циклов', 'Calendar years only (e.g. 2019–2024). No exact dates.', 'Только годы (например, 2019–2024). Точные даты не указываются.', 'text', 'O', example='2019-2024'),
            F('submission_scope', 'Scope', 'Охват', 'Whether every started cycle is reported (complete cycles) or a subset.', 'Сообщаются ли все начатые циклы (полные циклы) или только часть.', 'enum', 'R', allowed='submission_scope'),
            F('cycle_selection', 'How cycles were selected', 'Как отбирались циклы', 'Describe inclusion criteria; state explicitly whether cycles without transfer or without pregnancy are included.', 'Критерии включения; явно укажите, включены ли циклы без переноса и без беременности.', 'text', 'R'),
            F('overlap_with_published', 'Overlap with published data', 'Пересечение с опубликованными данными', 'Whether any cycle, embryo or child has already been reported elsewhere (to avoid double counting).', 'Сообщались ли какие-либо циклы, эмбрионы или дети ранее (чтобы не учесть их дважды).', 'enum', 'R', allowed='overlap'),
            F('overlap_details', 'Overlap details', 'Детали пересечения', 'DOI(s) and identifiers of the previously reported cases.', 'DOI и обозначения ранее описанных случаев.', 'text', 'C'),
            F('ethics_approval', 'Ethics approval', 'Одобрение этического комитета', 'Name of the committee and approval reference, or the reason approval was not required.', 'Название комитета и номер решения либо причина, по которой одобрение не требовалось.', 'text', 'R'),
            F('consent_for_research_use', 'Consent for research use', 'Согласие на использование в исследованиях', 'Patients consented to research use of de-identified data (yes / no / unknown).', 'Пациенты согласились на использование обезличенных данных в исследованиях.', 'enum', 'R', allowed='yes_no_unknown'),
            F('deidentification_confirmed', 'De-identification confirmed', 'Обезличивание подтверждено', 'The submitter confirms that the file contains no direct identifiers (names, exact dates, addresses, record numbers).', 'Автор подтверждает, что в файле нет прямых идентификаторов (имён, точных дат, адресов, номеров карт).', 'enum', 'R', allowed='yes_no'),
            F('authority_to_share', 'Authority to share', 'Полномочия на передачу', 'The submitter is authorised by the institution / data owner to share these data with the Atlas.', 'Автор уполномочен организацией или владельцем данных передать их в Atlas.', 'enum', 'R', allowed='yes_no'),
            F('permission_public_display', 'Permission for public display', 'Разрешение на публичный показ', 'After validation, aggregated and de-identified values may be shown in the Atlas with attribution.', 'После проверки обезличенные значения можно показывать в Atlas с указанием авторства.', 'enum', 'R', allowed='yes_no'),
            F('embargo_until', 'Embargo until', 'Эмбарго до', 'Year (or year-month) before which the data must not be shown publicly, if any.', 'Год (или год и месяц), до которого данные нельзя показывать публично, если есть ограничение.', 'text', 'O', example='2027-06'),
            F('preferred_credit', 'Preferred credit line', 'Предпочтительная форма указания вклада', 'How the contributing group wishes to be credited in the Atlas.', 'Как группа хочет, чтобы её вклад был указан в Atlas.', 'text', 'O'),
            F('notes', 'Notes', 'Примечания', 'Anything the reviewer should know.', 'Всё, что следует знать при проверке.', 'text', 'O'),
        ],
    ),
    Sheet(
        'Cycle', 'Cycle', 'Цикл',
        'One row per started IVF/PGT-mt cycle — including cycles that ended without a suitable embryo or without transfer.',
        'Одна строка на начатый цикл ЭКО/ПГТ-мт — включая циклы, завершившиеся без пригодного эмбриона или без переноса.',
        [
            F('cycle_id', 'Cycle ID', 'ID цикла', 'De-identified cycle code, unique within the file. Must not be derivable from clinic record numbers.', 'Обезличенный код цикла, уникальный в файле. Не должен выводиться из номеров медицинских карт.', 'id', 'R', example='C001'),
            F('study_id', 'Study ID', 'ID исследования', 'Link to the Study sheet.', 'Ссылка на лист Study.', 'ref', 'R', allowed='Study.study_id'),
            F('patient_code', 'Patient code', 'Код пациентки', 'De-identified code that links several cycles of the same patient. Not a name or initials.', 'Обезличенный код, связывающий несколько циклов одной пациентки. Не имя и не инициалы.', 'id', 'R', example='P01'),
            F('cycle_number', 'Cycle number for the patient', 'Номер цикла пациентки', 'Order of this cycle among the patient’s reported cycles (1, 2, …).', 'Порядковый номер цикла среди сообщаемых циклов пациентки.', 'integer', 'O', allowed='≥1'),
            F('age_group', 'Age group at retrieval', 'Возрастная группа при пункции', 'Patient age band at oocyte retrieval. Exact age or date of birth are not collected.', 'Возрастная группа пациентки на момент пункции. Точный возраст и дата рождения не собираются.', 'enum', 'R', allowed='age_group'),
            F('mtdna_variant', 'mtDNA variant', 'Вариант мтДНК', 'HGVS-style mtDNA variant (rCRS, NC_012920.1).', 'Вариант мтДНК в нотации HGVS (rCRS, NC_012920.1).', 'text', 'R', example='m.3243A>G'),
            F('gene', 'Gene', 'Ген', 'Gene affected by the variant.', 'Ген, в котором находится вариант.', 'text', 'O', example='MT-TL1'),
            F('variant_class', 'Variant classification', 'Классификация варианта', 'Classification at the time of testing; neutral polymorphisms are accepted as methodological controls.', 'Классификация на момент тестирования; нейтральные полиморфизмы принимаются как методический контроль.', 'enum', 'R', allowed='variant_class'),
            F('family_history', 'Family history', 'Семейный анамнез', 'How the variant came to attention in the family.', 'Как вариант был выявлен в семье.', 'enum', 'R', allowed='family_history'),
            F('family_phenotype', 'Phenotype in the family', 'Фенотип в семье', 'Syndrome or main clinical features in affected relatives, without identifying details.', 'Синдром или основные проявления у больных родственников без идентифицирующих подробностей.', 'text', 'O', example='MELAS in a previous child'),
            F('carrier_clinical_status', 'Carrier clinical status', 'Клинический статус носительницы', 'Clinical status of the patient (oocyte provider). Carrier tissue heteroplasmy goes to the Samples sheet.', 'Клинический статус пациентки. Гетероплазмия в тканях носительницы — на листе Samples.', 'enum', 'R', allowed='carrier_status'),
            F('fertilisation_method', 'Fertilisation method', 'Метод оплодотворения', 'ICSI is usual in PGT to avoid contamination; state if conventional IVF was used.', 'Обычно ИКСИ; укажите, если применялось стандартное ЭКО.', 'enum', 'R', allowed='fertilisation_method'),
            _count('oocytes_retrieved', 'Oocytes retrieved', 'Получено ооцитов', 'Cumulus–oocyte complexes retrieved.', 'Число полученных ооцит-кумулюсных комплексов.', ex='12'),
            _count('oocytes_mii', 'Mature (MII) oocytes', 'Зрелые ооциты (MII)', 'Metaphase II oocytes.', 'Ооциты метафазы II.'),
            _count('oocytes_immature', 'Immature oocytes (GV/MI)', 'Незрелые ооциты (GV/MI)', 'GV and MI oocytes.', 'Ооциты на стадиях GV и MI.', req='O'),
            _count('oocytes_degenerated', 'Degenerated oocytes', 'Дегенерировавшие ооциты', 'Oocytes degenerated after retrieval or injection.', 'Ооциты, дегенерировавшие после пункции или инъекции.', req='O'),
            _count('zygotes_2pn', 'Normally fertilised (2PN)', 'Нормально оплодотворённые (2PN)', 'Zygotes with two pronuclei.', 'Зиготы с двумя пронуклеусами.'),
            _count('zygotes_abnormal', 'Abnormally fertilised (0PN/1PN/3PN)', 'Аномально оплодотворённые (0PN/1PN/3PN)', 'Abnormally fertilised oocytes.', 'Аномально оплодотворённые ооциты.', req='O'),
            _count('oocytes_unfertilised', 'Unfertilised oocytes', 'Неоплодотворившиеся ооциты', 'Inseminated/injected oocytes that did not fertilise.', 'Ооциты, не оплодотворившиеся после инсеминации или ИКСИ.', req='O'),
            _count('embryos_cleavage_d3', 'Cleavage embryos on day 3', 'Дробящиеся эмбрионы на 3-и сутки', 'Embryos cleaving on day 3.', 'Эмбрионы, дробящиеся на 3-и сутки.', req='O'),
            _count('blastocysts_total', 'Blastocysts formed', 'Сформировалось бластоцист', 'Blastocysts on days 5–7.', 'Бластоцисты на 5–7-е сутки.'),
            _count('embryos_arrested', 'Embryos arrested', 'Остановившиеся эмбрионы', 'Embryos that arrested or degenerated before biopsy.', 'Эмбрионы, остановившиеся в развитии или дегенерировавшие до биопсии.', req='O'),
            _count('embryos_biopsied', 'Embryos biopsied', 'Эмбрионов биопсировано', 'Embryos (or oocytes, for polar body testing) with a biopsy sent for testing.', 'Эмбрионы (или ооциты при анализе полярных телец), биоптат которых направлен на исследование.'),
            _count('embryos_with_pgtmt_result', 'Embryos with a PGT-mt result', 'Эмбрионы с результатом ПГТ-мт', 'Biopsied embryos that received a conclusive PGT-mt result.', 'Биопсированные эмбрионы, получившие результат ПГТ-мт.'),
            _count('embryos_suitable', 'Embryos suitable for transfer', 'Эмбрионы, пригодные для переноса', 'Embryos meeting all transfer criteria of this cycle (heteroplasmy cut-off, PGT-A, morphology).', 'Эмбрионы, отвечающие всем критериям переноса в этом цикле (порог гетероплазмии, ПГТ-А, морфология).'),
            F('heteroplasmy_cutoff_pct', 'Heteroplasmy cut-off used', 'Применённый порог гетероплазмии', 'Cut-off used for the transfer decision in this cycle, if any.', 'Порог, применявшийся для решения о переносе в этом цикле, если был.', 'percent', 'C', '%', '0–100', '18'),
            F('cutoff_rationale', 'Rationale for the cut-off', 'Обоснование порога', 'Source or reasoning for the cut-off.', 'Источник или обоснование порога.', 'text', 'O'),
            F('pgta_performed', 'PGT-A performed', 'Проводился ПГТ-А', 'Whether PGT-A was performed on the same biopsies.', 'Проводился ли ПГТ-А на тех же биоптатах.', 'enum', 'R', allowed='yes_no_unknown'),
            F('cycle_endpoint', 'Cycle endpoint', 'Чем завершился цикл', 'How the cycle ended. Cycles without a suitable embryo are as important as cycles with a transfer.', 'Чем завершился цикл. Циклы без пригодного эмбриона так же важны, как циклы с переносом.', 'enum', 'R', allowed='cycle_endpoint'),
            F('notes', 'Notes', 'Примечания', 'Free text without identifying details.', 'Свободный текст без идентифицирующих подробностей.', 'text', 'O'),
        ],
    ),
    Sheet(
        'Embryos', 'Embryos', 'Эмбрионы',
        'One row per embryo (or tested oocyte) — every embryo of the cycle, not only the transferred ones.',
        'Одна строка на эмбрион (или исследованный ооцит) — все эмбрионы цикла, а не только перенесённые.',
        [
            F('embryo_id', 'Embryo ID', 'ID эмбриона', 'De-identified embryo code, unique within the file.', 'Обезличенный код эмбриона, уникальный в файле.', 'id', 'R', example='C001-E03'),
            F('cycle_id', 'Cycle ID', 'ID цикла', 'Link to the Cycle sheet.', 'Ссылка на лист Cycle.', 'ref', 'R', allowed='Cycle.cycle_id'),
            F('embryo_origin', 'Origin', 'Происхождение', 'Clinical PGT cycle or embryo donated for research.', 'Клинический цикл ПГТ или эмбрион, переданный для исследований.', 'enum', 'R', allowed='embryo_origin'),
            F('stage_at_biopsy', 'Stage at biopsy', 'Стадия при биопсии', 'Developmental stage when the sample was taken. Use “arrested” / “degenerated” for embryos analysed after arrest.', 'Стадия развития при взятии образца; для эмбрионов, исследованных после остановки, — arrested / degenerated.', 'enum', 'R', allowed='stage'),
            F('biopsy_day', 'Day of biopsy', 'День биопсии', 'Day after insemination (day 0).', 'Сутки после инсеминации (день 0).', 'integer', 'C', 'day', '0–7', '5'),
            F('highest_stage_reached', 'Highest stage reached', 'Максимальная достигнутая стадия', 'Most advanced stage reached in culture.', 'Самая поздняя стадия, достигнутая в культуре.', 'enum', 'O', allowed='stage'),
            F('morphology_system', 'Morphology system', 'Система морфологической оценки', 'Grading system used.', 'Использованная система оценки.', 'enum', 'C', allowed='morphology_system'),
            F('morphology_grade', 'Morphology grade', 'Морфологическая оценка', 'Grade as recorded in the stated system.', 'Оценка в указанной системе.', 'text', 'C', example='4AB'),
            F('pgta_result', 'PGT-A result', 'Результат ПГТ-А', 'Chromosomal result; “not_tested” if PGT-A was not done.', 'Хромосомный результат; not_tested, если ПГТ-А не проводился.', 'enum', 'R', allowed='pgta_result'),
            F('pgta_details', 'PGT-A details', 'Детали ПГТ-А', 'Karyotype summary, e.g. “+16”, “mosaic −7 (40%)”.', 'Кратко кариотип, например «+16», «мозаицизм −7 (40%)».', 'text', 'O'),
            F('decision_heteroplasmy_pct', 'Heteroplasmy used for the decision', 'Гетероплазмия, на основе которой принято решение', 'The single value used for the clinical decision (all individual measurements go to the Samples sheet).', 'Единственное значение, на основании которого принималось клиническое решение (все отдельные измерения — на листе Samples).', 'percent', 'C', '%', '0–100'),
            F('decision_value_qualifier', 'Qualifier of the decision value', 'Тип решающего значения', 'Distinguishes an exact value, a value below the detection limit, an interval, a failed analysis, etc.', 'Различает точное значение, значение ниже предела обнаружения, интервал, неудачный анализ и т. п.', 'enum', 'C', allowed='value_qualifier'),
            F('suitable_for_transfer', 'Suitable for transfer', 'Пригоден для переноса', 'Whether the embryo met all transfer criteria of the cycle.', 'Отвечал ли эмбрион всем критериям переноса в цикле.', 'enum', 'R', allowed='suitability'),
            F('transfer_status', 'Transfer status', 'Статус переноса', 'What happened to the embryo.', 'Что произошло с эмбрионом.', 'enum', 'R', allowed='transfer_status'),
            F('reason_not_transferred', 'Reason not transferred', 'Причина отказа от переноса', 'Main reason the embryo was not transferred; “not_applicable” if transferred.', 'Основная причина отказа от переноса; not_applicable, если эмбрион перенесён.', 'enum', 'C', allowed='not_transferred_reason'),
            F('outcome_id', 'Outcome ID', 'ID исхода', 'For transferred embryos: link to the Clinical_Outcomes sheet.', 'Для перенесённых эмбрионов: ссылка на лист Clinical_Outcomes.', 'ref', 'C', allowed='Clinical_Outcomes.outcome_id'),
            F('notes', 'Notes', 'Примечания', 'Free text without identifying details.', 'Свободный текст без идентифицирующих подробностей.', 'text', 'O'),
        ],
    ),
    Sheet(
        'Samples_Measurements', 'Samples and Measurements', 'Образцы и измерения',
        'One row per measurement: carrier tissues, oocytes, polar bodies, embryo fractions, prenatal and postnatal samples. Replicates are separate rows.',
        'Одна строка на измерение: ткани носительницы, ооциты, полярные тельца, фракции эмбриона, пренатальные и постнатальные образцы. Повторы — отдельными строками.',
        [
            F('measurement_id', 'Measurement ID', 'ID измерения', 'Unique code of this measurement.', 'Уникальный код измерения.', 'id', 'R', example='M0001'),
            F('sample_id', 'Sample ID', 'ID образца', 'Code of the physical sample; repeated measurements of one sample share it.', 'Код физического образца; повторные измерения одного образца имеют общий код.', 'id', 'R', example='S0001'),
            F('subject_level', 'Subject', 'Объект', 'Whose material this is.', 'Чей это материал.', 'enum', 'R', allowed='subject_level'),
            F('linked_id', 'Linked ID', 'Связанный ID', 'cycle_id for carrier tissues, oocytes and polar bodies; embryo_id for embryos; outcome_id for fetal or child samples.', 'cycle_id — для тканей носительницы, ооцитов и полярных телец; embryo_id — для эмбрионов; outcome_id — для образцов плода или ребёнка.', 'ref', 'R', allowed='Cycle.cycle_id | Embryos.embryo_id | Clinical_Outcomes.outcome_id'),
            F('sample_type', 'Sample type', 'Тип образца', 'Material analysed.', 'Исследованный материал.', 'enum', 'R', allowed='sample_type'),
            F('sample_type_other', 'Sample type (other)', 'Тип образца (другое)', 'Description if sample_type is “other”.', 'Описание, если выбран тип other.', 'text', 'C'),
            F('cells_in_sample', 'Approx. number of cells', 'Примерное число клеток', 'Approximate number of cells in the sample (e.g. 5 for a TE biopsy).', 'Примерное число клеток в образце (например, 5 для биоптата ТЭ).', 'integer', 'O', 'cells', '≥1'),
            F('timepoint_value', 'Time point', 'Срок', 'Embryo day, gestational week or child age, in the unit given in the next column.', 'Сутки развития эмбриона, неделя беременности или возраст ребёнка в единицах следующего столбца.', 'decimal', 'C'),
            F('timepoint_unit', 'Time-point unit', 'Единица срока', 'Unit of the time point.', 'Единица срока.', 'enum', 'C', allowed='timepoint_unit'),
            F('method_id', 'Method ID', 'ID метода', 'Link to the Methods sheet.', 'Ссылка на лист Methods.', 'ref', 'R', allowed='Methods.method_id'),
            F('measurement_kind', 'Measurement kind', 'Вид измерения', 'A single replicate, a mean of replicates, or a single measurement without replicates.', 'Отдельный повтор, среднее по повторам или однократное измерение.', 'enum', 'R', allowed='measurement_kind'),
            F('replicate_no', 'Replicate number', 'Номер повтора', 'For single_replicate rows: 1, 2, 3 …', 'Для строк single_replicate: 1, 2, 3 …', 'integer', 'C', allowed='≥1'),
            F('n_replicates', 'Number of replicates', 'Число повторов', 'For mean_of_replicates rows.', 'Для строк mean_of_replicates.', 'integer', 'C', allowed='≥1'),
            F('heteroplasmy_pct', 'Heteroplasmy', 'Гетероплазмия', 'Measured proportion of mtDNA carrying the variant. Leave blank unless value_qualifier is “exact”. 0 only if the method returned a quantified 0.', 'Измеренная доля мтДНК с вариантом. Заполняется только при value_qualifier = exact. 0 — только если метод выдал количественный ноль.', 'percent', 'C', '%', '0–100', '12.5'),
            F('value_qualifier', 'Value qualifier', 'Тип значения', 'Separates an exact value, a value below the detection limit, a published interval, a failed or not-performed analysis.', 'Различает точное значение, значение ниже предела обнаружения, опубликованный интервал, неудачный или непроведённый анализ.', 'enum', 'R', allowed='value_qualifier'),
            F('value_low_pct', 'Lower bound', 'Нижняя граница', 'For “interval” and “greater_than”.', 'Для interval и greater_than.', 'percent', 'C', '%', '0–100'),
            F('value_high_pct', 'Upper bound', 'Верхняя граница', 'For “interval” and “less_than”.', 'Для interval и less_than.', 'percent', 'C', '%', '0–100'),
            F('lod_pct', 'Limit of detection (sample)', 'Предел обнаружения (для образца)', 'Detection limit applying to this sample; required for “below_lod”.', 'Предел обнаружения для этого образца; обязателен при below_lod.', 'percent', 'C', '%', '0–100', '1'),
            F('sd_pct', 'SD between replicates', 'SD между повторами', 'Standard deviation for mean_of_replicates rows.', 'Стандартное отклонение для строк mean_of_replicates.', 'percent', 'O', 'pp', '≥0'),
            F('text_result', 'Qualitative result', 'Качественный результат', 'For “qualitative” rows: the result as reported (e.g. “homoplasmic wild type”).', 'Для qualitative: результат так, как он сообщён (например, «гомоплазмия дикого типа»).', 'text', 'C'),
            F('qc_passed', 'QC passed', 'Контроль качества пройден', 'Whether the measurement passed the laboratory QC criteria.', 'Прошло ли измерение лабораторный контроль качества.', 'enum', 'R', allowed='yes_no_unknown'),
            F('used_for_decision', 'Used for the clinical decision', 'Использовано для клинического решения', 'Whether this value informed the transfer decision.', 'Учитывалось ли значение при решении о переносе.', 'enum', 'O', allowed='yes_no_unknown'),
            F('mtdna_copy_number', 'mtDNA copy number', 'Копийность мтДНК', 'If measured; state the unit (per cell, relative).', 'Если измерялась; укажите единицу (на клетку, относительная).', 'text', 'O'),
            F('notes', 'Notes', 'Примечания', 'Free text without identifying details.', 'Свободный текст без идентифицирующих подробностей.', 'text', 'O'),
        ],
    ),
    Sheet(
        'Methods', 'Methods', 'Методы',
        'One row per analytical method or protocol version used in the measurements.',
        'Одна строка на аналитический метод или версию протокола, использованные в измерениях.',
        [
            F('method_id', 'Method ID', 'ID метода', 'Unique code of the method / protocol version.', 'Уникальный код метода или версии протокола.', 'id', 'R', example='NGS-v2'),
            F('study_id', 'Study ID', 'ID исследования', 'Link to the Study sheet.', 'Ссылка на лист Study.', 'ref', 'R', allowed='Study.study_id'),
            F('laboratory', 'Laboratory', 'Лаборатория', 'Laboratory that performed the analysis.', 'Лаборатория, выполнившая анализ.', 'text', 'R'),
            F('analytical_method', 'Analytical method', 'Аналитический метод', 'Principle of heteroplasmy quantification.', 'Принцип количественного определения гетероплазмии.', 'enum', 'R', allowed='analytical_method'),
            F('method_details', 'Method details', 'Детали метода', 'Platform, kit, chemistry, software and version.', 'Платформа, набор, химия, программа и версия.', 'text', 'R'),
            F('wga', 'Whole-genome amplification', 'Полногеномная амплификация', 'WGA performed before mtDNA analysis.', 'Проводилась ли полногеномная амплификация перед анализом мтДНК.', 'enum', 'R', allowed='wga'),
            F('wga_kit', 'WGA kit', 'Набор для WGA', 'Kit name and version.', 'Название и версия набора.', 'text', 'C'),
            F('target_region', 'Target region', 'Анализируемый участок', 'Amplicon or region of mtDNA analysed (rCRS coordinates), or “whole mtDNA”.', 'Ампликон или участок мтДНК (координаты rCRS) либо «вся мтДНК».', 'text', 'R', example='m.3150-3380'),
            F('read_depth_median', 'Median read depth', 'Медианная глубина прочтения', 'For sequencing methods: median depth at the variant position.', 'Для секвенирования: медианная глубина в позиции варианта.', 'integer', 'C', 'reads', '≥0'),
            F('read_depth_min_accepted', 'Minimum accepted depth', 'Минимальная допустимая глубина', 'Depth below which a result is rejected.', 'Глубина, ниже которой результат отклоняется.', 'integer', 'C', 'reads', '≥0'),
            F('lod_pct', 'Limit of detection', 'Предел обнаружения', 'Validated limit of detection of the method.', 'Валидированный предел обнаружения метода.', 'percent', 'R', '%', '0–100', '1'),
            F('loq_pct', 'Limit of quantification', 'Предел количественного определения', 'Lowest level quantified with stated precision.', 'Наименьший уровень, определяемый с заявленной точностью.', 'percent', 'O', '%', '0–100'),
            F('linearity_check', 'Linearity check', 'Проверка линейности', 'How linearity was verified (e.g. mixtures of known heteroplasmy, range, R²).', 'Как проверяли линейность (например, смеси с известной гетероплазмией, диапазон, R²).', 'text', 'R'),
            F('precision_pp', 'Technical precision', 'Техническая точность', 'Typical technical variation between replicates.', 'Типичный технический разброс между повторами.', 'decimal', 'O', 'pp', '≥0'),
            F('positive_controls', 'Positive controls', 'Положительные контроли', 'Positive controls used.', 'Использованные положительные контроли.', 'text', 'R'),
            F('negative_controls', 'Negative / blank controls', 'Отрицательные и холостые контроли', 'Negative and no-template controls; contamination monitoring.', 'Отрицательные контроли и контроли без матрицы; контроль контаминации.', 'text', 'R'),
            F('ado_assessment', 'Amplification bias / ADO assessment', 'Оценка смещения амплификации / ADO', 'How amplification bias or allele drop-out was assessed in small samples.', 'Как оценивали смещение амплификации или выпадение аллеля в малых образцах.', 'text', 'O'),
            F('validation_reference', 'Validation reference', 'Ссылка на валидацию', 'DOI or document describing method validation.', 'DOI или документ с описанием валидации метода.', 'text', 'O'),
            F('guidelines_followed', 'Guidelines followed', 'Применённые рекомендации', 'Professional guidelines for genetic testing in mitochondrial disease that the protocol follows.', 'Профессиональные рекомендации по генетическому тестированию при митохондриальных заболеваниях, которым следует протокол.', 'text', 'O'),
            F('accreditation', 'Accreditation', 'Аккредитация', 'Laboratory accreditation (e.g. ISO 15189), if any.', 'Аккредитация лаборатории (например, ISO 15189), если есть.', 'text', 'O'),
            F('pgta_method', 'PGT-A method', 'Метод ПГТ-А', 'Platform used for PGT-A, if performed.', 'Платформа ПГТ-А, если он проводился.', 'text', 'O'),
            F('notes', 'Notes', 'Примечания', 'Free text.', 'Свободный текст.', 'text', 'O'),
        ],
    ),
    Sheet(
        'Clinical_Outcomes', 'Clinical Outcomes', 'Клинические исходы',
        'One row per outcome event: cycles without a suitable embryo, transfers without implantation, pregnancy losses, births and each follow-up visit.',
        'Одна строка на событие: цикл без пригодного эмбриона, перенос без имплантации, потеря беременности, роды и каждый визит наблюдения.',
        [
            F('outcome_id', 'Outcome ID', 'ID исхода', 'Unique code of the outcome event.', 'Уникальный код события.', 'id', 'R', example='O001'),
            F('cycle_id', 'Cycle ID', 'ID цикла', 'Link to the Cycle sheet.', 'Ссылка на лист Cycle.', 'ref', 'R', allowed='Cycle.cycle_id'),
            F('embryo_ids', 'Transferred embryo IDs', 'ID перенесённых эмбрионов', 'Embryo ID(s) transferred, separated by “;”. Blank for cycles without transfer.', 'ID перенесённых эмбрионов через «;». Пусто для циклов без переноса.', 'ref', 'C', allowed='Embryos.embryo_id'),
            F('outcome_event', 'Event', 'Событие', 'Outcome event. Report negative outcomes with the same care as live births.', 'Событие исхода. Отрицательные исходы сообщаются так же подробно, как живорождения.', 'enum', 'R', allowed='outcome_event'),
            F('gestational_age_weeks', 'Gestational age at event', 'Срок беременности при событии', 'Completed weeks of gestation, if applicable.', 'Полные недели беременности, если применимо.', 'decimal', 'C', 'weeks', '0–44'),
            F('n_fetuses', 'Number of fetuses', 'Число плодов', 'For clinical pregnancies.', 'Для клинических беременностей.', 'integer', 'C', allowed='≥0'),
            F('pregnancy_loss_cause', 'Cause of loss (if known)', 'Причина потери (если известна)', 'E.g. fetal karyotype result; no identifying details.', 'Например, результат кариотипирования плода; без идентифицирующих подробностей.', 'text', 'O'),
            F('prenatal_testing', 'Prenatal testing', 'Пренатальная диагностика', 'Invasive prenatal testing performed; heteroplasmy values go to the Samples sheet.', 'Проведённая инвазивная пренатальная диагностика; значения гетероплазмии — на листе Samples.', 'enum', 'C', allowed='prenatal_test'),
            F('child_code', 'Child code', 'Код ребёнка', 'De-identified code of the child; links follow-up visits of one child.', 'Обезличенный код ребёнка; связывает визиты наблюдения одного ребёнка.', 'id', 'C', example='K01'),
            F('child_age_months', 'Child age at assessment', 'Возраст ребёнка при оценке', 'Age in months at this assessment. Exact birth dates are not collected.', 'Возраст в месяцах при этой оценке. Точные даты рождения не собираются.', 'decimal', 'C', 'months', '≥0'),
            F('child_status', 'Child clinical status', 'Клинический статус ребёнка', 'Clinical status at this assessment.', 'Клинический статус при этой оценке.', 'enum', 'C', allowed='child_status'),
            F('assessment_details', 'Assessment details', 'Детали оценки', 'What was assessed (e.g. neurological examination, lactate), without identifying details.', 'Что оценивалось (например, неврологический осмотр, лактат), без идентифицирующих подробностей.', 'text', 'O'),
            F('notes', 'Notes', 'Примечания', 'Free text without identifying details.', 'Свободный текст без идентифицирующих подробностей.', 'text', 'O'),
        ],
    ),
]

REQ_TEXT = {
    'R': ('required', 'обязательно'),
    'C': ('conditional', 'при условии'),
    'O': ('if available', 'по возможности'),
}

# ---------------------------------------------------------------- инструкции

INSTRUCTIONS_EN = [
    ('title', f'PGT-mt Research Data Submission · template v{VERSION} · {RELEASED}'),
    ('sub', 'PGT-mt Atlas — an independent, continuously updated scientific database. Created and curated by Ilya Mazunin.'),
    ('motto', 'Every outcome matters. Negative results are data, too.'),
    ('h', 'Why this template'),
    ('p', 'Successful PGT-mt cycles are published more often than cycles without a suitable embryo, without implantation or with a pregnancy loss. The literature therefore over-represents favourable outcomes. Complete cycles — every oocyte, every embryo and every outcome — let the Atlas estimate how mtDNA variants are distributed and how often PGT-mt leads to a transfer and a healthy child.'),
    ('h', 'Sheets and how they link'),
    ('p', 'Study (study_id) → Cycle (cycle_id) → Embryos (embryo_id) → Samples_Measurements (measurement_id). Methods (method_id) are referenced from each measurement; Clinical_Outcomes (outcome_id) are linked to cycles and transferred embryos. Data_Dictionary defines every field; Lists holds the allowed values used by the drop-down menus.'),
    ('p', 'Row 1 of each data sheet holds machine-readable field codes, row 2 a short label and whether the field is required. Enter data from row 3. Hover over a field code to read its definition.'),
    ('h', 'Zero, missing and below the detection limit are different'),
    ('li', 'A blank cell means “not reported / unknown”. Never type 0 for unknown.'),
    ('li', '0 in a count means “counted, none” (e.g. 0 blastocysts).'),
    ('li', 'A heteroplasmy value below the limit of detection: value_qualifier = below_lod, heteroplasmy_pct blank, lod_pct filled. Do not type 0.'),
    ('li', 'Only an interval or a bound was published: value_qualifier = interval / less_than / greater_than with the bounds; no midpoint is imputed.'),
    ('li', 'An analysis that failed QC or was not performed: value_qualifier = failed_qc / not_measured. These rows are kept — they are part of the cycle.'),
    ('h', 'Never include direct identifiers'),
    ('li', 'No names, initials, exact dates of birth, retrieval, transfer or delivery, addresses, phone numbers, e-mails of patients, medical record or biobank numbers, photographs.'),
    ('li', 'Use age groups, years and relative time points (embryo day, gestational week, child age in months).'),
    ('li', 'Use codes (P01, C001, C001-E03, K01) that cannot be traced back to clinic records outside your institution.'),
    ('h', 'What happens after submission'),
    ('p', 'Submitted → Under review → Validated → Included or Rejected. Nothing is added to public charts automatically. Each submission is checked for completeness, internal consistency, overlap with published data, authority to share, permission for use, de-identification and ethics. The evidence status (peer-reviewed / published without peer review / unpublished) stays visible next to every value.'),
    ('h', 'Credit'),
    ('p', 'Every included value keeps its source: the study, its authors and DOI or the contributing group. Contributing groups may be invited to collaborate on joint scientific publications; authorship is decided case by case according to the ICMJE criteria and is not automatic.'),
    ('h', 'How to send'),
    ('p', 'Before sending any file, contact the curator through the Contribute Data page of the Atlas to agree on a secure transfer channel. Do not send patient files by unencrypted e-mail.'),
]

INSTRUCTIONS_RU = [
    ('title', f'PGT-mt Research Data Submission · шаблон v{VERSION} · {RELEASED}'),
    ('sub', 'PGT-mt Atlas — независимая, постоянно обновляемая научная база данных. Создатель и куратор — Илья Мазунин.'),
    ('motto', 'Важен каждый исход. Отрицательный результат — тоже данные.'),
    ('h', 'Зачем нужен шаблон'),
    ('p', 'Успешные циклы ПГТ-мт публикуются чаще, чем циклы без пригодного эмбриона, без имплантации или с потерей беременности, поэтому литература смещена в сторону благоприятных исходов. Полные циклы — каждый ооцит, каждый эмбрион и каждый исход — позволяют оценить распределение вариантов мтДНК и то, как часто ПГТ-мт приводит к переносу и рождению здорового ребёнка.'),
    ('h', 'Листы и их связи'),
    ('p', 'Study (study_id) → Cycle (cycle_id) → Embryos (embryo_id) → Samples_Measurements (measurement_id). Каждое измерение ссылается на метод (Methods, method_id); исходы (Clinical_Outcomes, outcome_id) связаны с циклами и перенесёнными эмбрионами. Data_Dictionary определяет каждое поле; Lists содержит допустимые значения для выпадающих списков.'),
    ('p', 'В строке 1 каждого листа данных — машиночитаемые коды полей, в строке 2 — краткое название и обязательность. Данные вводятся с 3-й строки. Определение поля показывается при наведении на его код.'),
    ('h', 'Ноль, пропуск и «ниже предела обнаружения» — разные значения'),
    ('li', 'Пустая ячейка — «не сообщается / неизвестно». Не ставьте 0 вместо «неизвестно».'),
    ('li', '0 в поле-счётчике — «посчитано, таких нет» (например, 0 бластоцист).'),
    ('li', 'Гетероплазмия ниже предела обнаружения: value_qualifier = below_lod, heteroplasmy_pct пусто, lod_pct заполнен. Не ставьте 0.'),
    ('li', 'Опубликован только интервал или граница: value_qualifier = interval / less_than / greater_than с границами; середина интервала не подставляется.'),
    ('li', 'Анализ не прошёл контроль качества или не проводился: value_qualifier = failed_qc / not_measured. Такие строки сохраняются — они часть цикла.'),
    ('h', 'Не включайте прямые идентификаторы'),
    ('li', 'Никаких имён, инициалов, точных дат рождения, пункции, переноса или родов, адресов, телефонов и e-mail пациентов, номеров карт или биобанка, фотографий.'),
    ('li', 'Используйте возрастные группы, годы и относительные сроки (сутки эмбриона, неделя беременности, возраст ребёнка в месяцах).'),
    ('li', 'Используйте коды (P01, C001, C001-E03, K01), по которым нельзя восстановить медицинские записи вне вашей организации.'),
    ('h', 'Что происходит после передачи'),
    ('p', 'Получено → На проверке → Проверено → Включено или Отклонено. Ничто не попадает в публичные графики автоматически. Каждое предложение проверяется на полноту, внутреннюю согласованность, пересечение с опубликованными данными, полномочия на передачу, разрешение на использование, обезличивание и соблюдение этических требований. Статус доказательности (рецензированная публикация / публикация без рецензирования / неопубликованные данные) остаётся виден рядом с каждым значением.'),
    ('h', 'Признание вклада'),
    ('p', 'У каждого включённого значения сохраняется источник: исследование, его авторы и DOI или группа, передавшая данные. Группы могут быть приглашены к совместным научным публикациям; авторство определяется в каждом случае отдельно по критериям ICMJE и не предоставляется автоматически.'),
    ('h', 'Как передать'),
    ('p', 'Прежде чем отправлять файл, свяжитесь с куратором через страницу «Поделиться данными» в Atlas и согласуйте защищённый канал передачи. Не отправляйте файлы с данными пациентов незашифрованной почтой.'),
]


# ---------------------------------------------------------------- сборка

THIN = Side(style='thin', color=RULE)
INK_LINE = Side(style='medium', color=INK)


def _style_instructions(ws, blocks, col):
    styles = {
        'title': Font(name='Georgia', size=16, bold=True, color=INK),
        'sub': Font(name='Georgia', size=10, italic=True, color=MUTED),
        'motto': Font(name='Georgia', size=13, italic=True, color=RED),
        'h': Font(name='Georgia', size=12, bold=True, color=INK),
        'p': Font(name='Calibri', size=11, color=INK2),
        'li': Font(name='Calibri', size=11, color=INK2),
    }
    row = 1
    for kind, text in blocks:
        if kind == 'h':
            row += 1
        c = ws.cell(row=row, column=col, value=('•  ' + text) if kind == 'li' else text)
        c.font = styles[kind]
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if kind == 'h':
            c.border = Border(bottom=THIN)
        row += 1
    return row


def build_workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Instructions'
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 96
    ws.column_dimensions['B'].width = 4
    ws.column_dimensions['C'].width = 96
    _style_instructions(ws, INSTRUCTIONS_EN, 1)
    _style_instructions(ws, INSTRUCTIONS_RU, 3)

    # допустимые значения — на листе Lists (создаётся последним), выпадающие списки ссылаются на именованные диапазоны
    list_ref = {name: f'list_{name}' for name in LISTS}

    head_font = Font(name='Calibri', size=11, bold=True, color=INK)
    label_font = Font(name='Calibri', size=9, italic=True, color=MUTED)
    req_font = Font(name='Calibri', size=9, italic=True, color=RED)
    last = FIRST_DATA_ROW + DATA_ROWS - 1
    for sh in SHEETS:
        ws = wb.create_sheet(sh.name)
        ws.sheet_properties.tabColor = INK
        ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=2)
        for j, f in enumerate(sh.fields, start=1):
            col = get_column_letter(j)
            c = ws.cell(row=1, column=j, value=f.code)
            c.font = head_font
            c.border = Border(bottom=INK_LINE)
            c.alignment = Alignment(vertical='bottom')
            note = f'{f.label_en} / {f.label_ru}\n\n{f.desc_en}\n\n{f.desc_ru}'
            if f.unit:
                note += f'\n\nUnit / единица: {f.unit}'
            if f.type == 'enum':
                note += f'\n\nAllowed values: see Lists → {f.allowed}'
            c.comment = Comment(note, 'PGT-mt Atlas', width=360, height=260)
            en, ru = REQ_TEXT[f.req]
            lab = ws.cell(row=2, column=j, value=f'{f.label_en} · {f.label_ru} — {en} / {ru}')
            lab.font = req_font if f.req == 'R' else label_font
            lab.alignment = Alignment(wrap_text=True, vertical='top')
            lab.border = Border(bottom=THIN)
            ws.column_dimensions[col].width = max(16, min(34, len(f.code) + 6))
            rng = f'{col}{FIRST_DATA_ROW}:{col}{last}'
            dv = None
            if f.type == 'enum':
                dv = DataValidation(type='list', formula1=f'={list_ref[f.allowed]}', allow_blank=True)
                dv.error = f'Use a value from Lists → {f.allowed} / Выберите значение из списка'
            elif f.type == 'percent':
                dv = DataValidation(type='decimal', operator='between', formula1='0', formula2='100', allow_blank=True)
                dv.error = 'Percent 0–100 (blank if not reported; use value_qualifier for below LOD) / Процент 0–100'
            elif f.type == 'integer':
                lo = '1' if f.allowed == '≥1' else '0'
                dv = DataValidation(type='whole', operator='greaterThanOrEqual', formula1=lo, allow_blank=True)
                dv.error = f'Whole number ≥{lo}; leave blank if unknown / Целое число ≥{lo}; пусто, если неизвестно'
            elif f.type == 'decimal':
                dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True)
                dv.error = 'Number ≥0 / Число ≥0'
            if dv is not None:
                dv.errorTitle = 'PGT-mt Atlas'
                dv.showErrorMessage = True
                dv.add(rng)
                ws.add_data_validation(dv)
        ws.row_dimensions[2].height = 42

    # Data_Dictionary
    dd = wb.create_sheet('Data_Dictionary')
    cols = ['sheet', 'field', 'label_en', 'label_ru', 'requirement', 'type', 'unit', 'allowed_values', 'definition_en', 'definition_ru', 'example']
    widths = [20, 28, 30, 30, 14, 10, 9, 40, 70, 70, 18]
    for j, (name, w) in enumerate(zip(cols, widths), start=1):
        c = dd.cell(row=1, column=j, value=name)
        c.font = head_font
        c.border = Border(bottom=INK_LINE)
        dd.column_dimensions[get_column_letter(j)].width = w
    r = 2
    for sh in SHEETS:
        for f in sh.fields:
            allowed = ', '.join(LISTS[f.allowed]) if f.type == 'enum' else f.allowed
            vals = [sh.name, f.code, f.label_en, f.label_ru, REQ_TEXT[f.req][0], f.type, f.unit, allowed, f.desc_en, f.desc_ru, f.example]
            for j, v in enumerate(vals, start=1):
                c = dd.cell(row=r, column=j, value=v or None)
                c.alignment = Alignment(wrap_text=True, vertical='top')
                c.border = Border(bottom=THIN)
            r += 1
    dd.freeze_panes = 'C2'
    dd.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{r - 1}'

    lists_ws = wb.create_sheet('Lists')
    for j, (name, values) in enumerate(LISTS.items(), start=1):
        col = get_column_letter(j)
        h = lists_ws.cell(row=1, column=j, value=name)
        h.font = head_font
        h.border = Border(bottom=INK_LINE)
        for i, v in enumerate(values, start=2):
            lists_ws.cell(row=i, column=j, value=v)
        lists_ws.column_dimensions[col].width = max(14, max(len(v) for v in values) + 2)
        wb.defined_names[list_ref[name]] = DefinedName(list_ref[name], attr_text=f'Lists!${col}$2:${col}${len(values) + 1}')
    lists_ws.freeze_panes = 'A2'

    wb.properties.creator = 'PGT-mt Atlas'
    wb.properties.lastModifiedBy = 'PGT-mt Atlas'
    wb.properties.title = f'PGT-mt Research Data Submission v{VERSION}'
    wb.properties.subject = 'Template for contributing PGT-mt research data to PGT-mt Atlas'
    fixed = dt.datetime(2026, 10, 10)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    return wb


def workbook_bytes() -> bytes:
    """xlsx с фиксированными метками времени — побайтно воспроизводимый."""
    raw = io.BytesIO()
    build_workbook().save(raw)
    src = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == 'docProps/core.xml':
                # openpyxl подставляет текущее время в modified при сохранении
                data = re.sub(rb'(<dcterms:modified[^>]*>)[^<]*(<)', rb'\g<1>2026-10-10T00:00:00Z\g<2>', data)
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 10, 10, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            dst.writestr(zi, data)
    return out.getvalue()


def summary() -> dict:
    return {
        'version': VERSION,
        'released': RELEASED,
        'file': f'templates/{FILENAME}',
        'sheets': [
            {
                'name': sh.name,
                'title_en': sh.title_en,
                'title_ru': sh.title_ru,
                'purpose_en': sh.purpose_en,
                'purpose_ru': sh.purpose_ru,
                'n_fields': len(sh.fields),
                'n_required': sum(f.req == 'R' for f in sh.fields),
            }
            for sh in SHEETS
        ],
        'value_qualifiers': LISTS['value_qualifier'],
        'evidence_tiers': LISTS['evidence_tier'],
    }


def summary_text() -> str:
    return json.dumps(summary(), ensure_ascii=False, indent=2) + '\n'


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--check', action='store_true', help='only compare with the committed files')
    args = ap.parse_args(argv)
    data, js = workbook_bytes(), summary_text()
    if args.check:
        stale = [p for p, new in ((OUT_XLSX, data), (OUT_JSON, js.encode())) if not p.exists() or p.read_bytes() != new]
        for p in stale:
            print(f'stale: {p.relative_to(ROOT)} — run python -m atlas_etl.submission_template', file=sys.stderr)
        return 1 if stale else 0
    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    OUT_XLSX.write_bytes(data)
    OUT_JSON.write_text(js, encoding='utf-8')
    print(f'{OUT_XLSX.relative_to(ROOT)} ({len(data) / 1024:.1f} KB), {OUT_JSON.relative_to(ROOT)}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
