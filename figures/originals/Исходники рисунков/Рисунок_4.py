"""Рисунок 4 — однородность гетероплазмии внутри эмбриона.
Данные: лист «Рисунок 4» (панель А: значения фракций 40 бластоцист) и лист «Бластомеры (дробление)»
(22 эмбриона стадии дробления, Monnot et al. 2011) сопроводительной таблицы.
Запуск: python3 Рисунок_4.py <Сопроводительная информация.xlsx> [выход.tif]
"""
import sys, math, statistics
import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D

XLSX = sys.argv[1]
OUT = sys.argv[2] if len(sys.argv) > 2 else 'Рисунок_4.tif'
wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb['Рисунок 4']
FR = ['Биоптат ТЭ', 'ВКМ', 'Оставшаяся ТЭ']
emb = {}
for r in ws.iter_rows(min_row=4, values_only=True):
    if r[0] is None or r[3] not in FR or not isinstance(r[4], (int, float)):
        if r[0] and str(r[0]).startswith('Панель Б'):
            break
        continue
    emb.setdefault(r[0], {'var': r[2], 'fr': {}})['fr'][r[3]] = float(r[4])
assert len(emb) == 40, len(emb)
wsb = wb['Бластомеры (дробление)']
cleav = [float(r[5]) for r in wsb.iter_rows(min_row=4, values_only=True)
         if r[0] and str(r[0]).startswith('Monnot') and isinstance(r[5], (int, float))]
assert len(cleav) == 22, len(cleav)

is_hv2 = lambda e: emb[e]['var'] == 'HV2 polymorphism'
PAT = dict(marker='o', ms=4.3, mfc='black', mec='black', mew=0.8)
HV = dict(marker='o', ms=4.3, mfc='white', mec='#808080', mew=1.0)

plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 9})
fig = plt.figure(figsize=(9.64, 4.78))
axA = fig.add_axes([0.065, 0.25, 0.24, 0.69])
axB = fig.add_axes([0.40, 0.25, 0.215, 0.69])
axC = fig.add_axes([0.71, 0.25, 0.28, 0.69])


def swarm(ax, xc, values, step, min_px, style, zorder=4):
    """Раскладывает совпадающие значения по горизонтали (детерминированно)."""
    fig.canvas.draw()
    placed = []
    for v in values:
        for k in [0] + [s * j for j in range(1, 40) for s in (1, -1)]:
            x = xc + k * step
            p = ax.transData.transform((x, v))
            if all(math.hypot(p[0] - q[0], p[1] - q[1]) >= min_px for q in placed):
                placed.append(p); break
        ax.plot([x], [v], ls='none', zorder=zorder, **style)


# ---------- А ----------
axA.axhspan(18, 30, color='#e3ecf7', lw=0, zorder=0)
axA.text(2.28, 24, '18–30%', color='#1565c0', ha='right', va='center', fontsize=8.5)
for e in sorted(emb, key=is_hv2, reverse=True):  # HV2 drawn first (under)
    xs = [i for i, f in enumerate(FR) if f in emb[e]['fr']]
    ys = [emb[e]['fr'][FR[i]] for i in xs]
    st = HV if is_hv2(e) else PAT
    axA.plot(xs, ys, color='#9e9e9e' if is_hv2(e) else '#444444', lw=0.9, zorder=2 if is_hv2(e) else 3, **st)
axA.set_xlim(-0.3, 2.3); axA.set_ylim(-3, 103)
axA.set_xticks([0, 1, 2]); axA.set_xticklabels(['Биоптат\nТЭ', 'ВКМ', 'Оставшаяся\nТЭ'])
axA.set_ylabel('Гетероплазмия, %')
n_p = sum(not is_hv2(e) for e in emb); n_h = sum(is_hv2(e) for e in emb)
axA.legend(handles=[Line2D([], [], color='#444444', lw=0.9, label=f'патогенные варианты (n = {n_p})', **PAT),
                    Line2D([], [], color='#9e9e9e', lw=0.9, label=f'нейтральный HV2 (n = {n_h})', **HV)],
           loc='upper left', bbox_to_anchor=(-0.05, -0.20), frameon=False, fontsize=8.5, handlelength=1.6)

# ---------- Б ----------
axB.axhspan(-5, 5, color='#eeeeee', lw=0, zorder=0)
axB.axhline(0, ls='--', color='#666666', lw=1.0, zorder=1)
groups = [('Оставшаяся ТЭ', False), ('Оставшаяся ТЭ', True), ('ВКМ', False), ('ВКМ', True)]
axB.set_xlim(-0.5, 3.5); axB.set_ylim(-12.5, 12.5)
for gx, (frac, hv) in enumerate(groups):
    d = sorted(emb[e]['fr'][frac] - emb[e]['fr']['Биоптат ТЭ'] for e in emb
               if is_hv2(e) == hv and frac in emb[e]['fr'])
    swarm(axB, gx, d, 0.05, 4.6 / 72 * fig.dpi * 1.02, HV if hv else PAT)
    axB.text(gx, -11, f'n = {len(d)}', ha='center', va='center', fontsize=8.5, color='#555555')
axB.set_xticks([0, 1, 2, 3]); axB.set_xticklabels(['патог.', 'HV2', 'патог.', 'HV2'])
axB.text(0.5, -0.12, 'оставшаяся ТЭ\n− биоптат', transform=axB.get_xaxis_transform(), ha='center', va='top', fontsize=9)
axB.text(2.5, -0.12, 'ВКМ\n− биоптат', transform=axB.get_xaxis_transform(), ha='center', va='top', fontsize=9)
axB.set_ylabel('Δ гетероплазмии, п.п.')
axB.text(0.62, 0.97, 'серая полоса — ±5 п.п.', transform=axB.transAxes, ha='center', va='top', fontsize=8.5, color='#555555')

# ---------- В ----------
rng = lambda e: max(emb[e]['fr'].values()) - min(emb[e]['fr'].values())
cols = [('Дробление,\n2 бластомера', sorted(cleav), PAT),
        ('Бластоциста,\nпатогенные', sorted(rng(e) for e in emb if not is_hv2(e)), PAT),
        ('Бластоциста,\nHV2', sorted(rng(e) for e in emb if is_hv2(e)), HV)]
axC.axhspan(0, 5, color='#eeeeee', lw=0, zorder=0)
axC.set_xlim(-0.5, 2.5); axC.set_ylim(-2.2, 11)
for gx, (lab, vals, st) in enumerate(cols):
    swarm(axC, gx, vals, 0.04, 4.6 / 72 * fig.dpi * 1.02, st)
    med = statistics.median(vals)
    axC.plot([gx - 0.23, gx + 0.23], [med, med], color='#d32f2f', lw=2.4, zorder=5, solid_capstyle='butt')
    axC.text(gx + 0.23, med + 0.55, f'{med:.1f}', color='#d32f2f', fontsize=9.5, ha='left', va='center')
    axC.text(gx, -1.6, f'n = {len(vals)}', ha='center', va='center', fontsize=8.5, color='#555555')
axC.set_xticks([0, 1, 2]); axC.set_xticklabels([c[0] for c in cols], fontsize=8.5)
axC.set_ylabel('Размах внутри эмбриона, п.п.')
axC.text(0.99, 0.97, 'красная черта — медиана', transform=axC.transAxes, ha='right', va='top', fontsize=8.5, color='#d32f2f')

for ax, letter in ((axA, 'А'), (axB, 'Б'), (axC, 'В')):
    for s in ('top', 'right'):
        ax.spines[s].set_visible(False)
    ax.text(0.0, 1.015, letter, transform=ax.transAxes, fontsize=11, fontweight='bold', va='bottom')
fig.savefig(OUT, dpi=600, bbox_inches='tight', pad_inches=0.05,
            **({'pil_kwargs': {'compression': 'tiff_lzw'}} if OUT.endswith('.tif') else {}))
print('saved', OUT, 'medians:', [round(statistics.median(c[1]), 2) for c in cols])
